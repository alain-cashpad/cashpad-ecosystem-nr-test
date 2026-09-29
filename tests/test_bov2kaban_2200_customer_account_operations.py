from __future__ import annotations

"""
Test de non-régression généré par /cp-test — ticket BOV2KABAN-2200
[BACK] Add "Customer Account Operations" predefined export (xlsx)

Fige le comportement OBSERVÉ et VALIDÉ le 2026-09-24 sur le staging
(host allowlisté : staging.cashpad.app), export prédéfini
`customer_account_operations` du site 4652 sur le 18/06/2024 (archive 25) :

    POST {STAGING_BASE_URL}/p/exports/api/1/site/{SITE}/predefined-exports
         → 201 {exportId}
    GET  {STAGING_BASE_URL}/p/exports/api/1/site/{SITE}/predefined-exports?exportId=…
         → [{id, status, bucketKey, …}]   (génération asynchrone, on poll)
    GET  {STAGING_BASE_URL}/p/exports/api/1/site/{SITE}/download-export/{historyId}
         → "<URL présignée>"               (chaîne JSON)
    GET  <URL présignée>                   → le xlsx

Même routage que `test_bov2kaban_2218_customers_csv_export.py` : préfixe `/p/`
obligatoire et **JWT sso** via `POST /p/sso/public/1/sign-in`.

⚠️ **Ce fichier ÉCRIT sur le staging** : chaque exécution crée une entrée
d'export prédéfini (et un objet dans le bucket) sur le site 4652. Aucune donnée
métier n'est touchée.

## Le jeu de données — pourquoi ce site et ce jour

Le 18/06/2024 est le seul jour du site 4652 portant des opérations de compte
**liées à un ticket** : toutes appartiennent au compte « Cashpad NR auto test »
(`ccc857d8-…`), les ~92 autres comptes du site sont des seeds sans ticket.

| Ticket | Opération (millièmes) | Total ticket | Ratio | Ligne export (TTC) |
|---|---|---|---|---|
| 570 | -13500 | 13.50 | 1   | +13.50 |
| 571 | +67500 | 27.00 | 2.5 | -67.50 |
| 573 | +13500 | 13.50 | 1   | -13.50 |
| —   | +200000 | — | — | Crédit 200 |

Le signe de l'export est l'opposé de celui de l'opération (= signe du paiement
« Compte client » sur le ticket). Le 571 est le seul cas où la proratisation
n'est pas triviale — c'est lui qui porte l'AC de proratisation. Aucun ratio < 1
n'existe dans la donnée.

## Ce que ce fichier ne fige pas

* **Les noms d'onglets** (`Résumé` / `Détail` / `Détail tickets`) : écart signalé
  avec BOV1 (`Summary` pour le premier). On fige leur nombre et leurs colonnes.
* **Le HT de la ligne d'opération du 571 à l'arrondi près** : au 2026-09-24 elle
  vaut -61.38 (base arrondie 24.55 × 2.5) alors que ses lignes produit somment
  -61.36 (base exacte 24.5454…). Écart signalé, non tranché : tolérance 0.02.
* **Le nom du site en A1** : on fige qu'il est présent, pas sa valeur.
* **Les opérations à montant nul** : absentes de la donnée, AC non exercée.

Le corps POST rejoue à l'identique celui du dashboard au 2026-09-24, y compris
ses deux bizarreries côté front (ticket FRONT séparé) : `format: "csv"` alors
que le back rend un xlsx, et `to == from` alors que l'archive 25 ferme à 16:01Z.
Le back s'appuie sur `cpFrom`/`cpTo` et les ignore.

Lancer :
    uv run --with pytest --with httpx --with python-dotenv --with openpyxl \
      pytest tests/test_bov2kaban_2200_customer_account_operations.py -v
"""

import io
import os
import time
from urllib.parse import urlparse

import httpx
import openpyxl
import pytest
from dotenv import load_dotenv

load_dotenv()


# ── Garde-fou anti-prod : la vérification ne cible que le staging ──
HOST_ALLOWLIST = {"staging.cashpad.app"}
# Le fichier est servi par le bucket de staging via une URL présignée.
STORAGE_HOST_ALLOWLIST = {"staging-object-storage.ams3.digitaloceanspaces.com"}

SITE = 4652

EXPORT_BODY = {
    "filters": {
        "computedTimeRanges": {
            "cpType": "archiveRanges",
            "cpFrom": 25,
            "cpTo": 25,
            "from": "2024-06-18T09:32:39.000Z",
            "to": "2024-06-18T09:32:39.000Z",
            "timezone": "Europe/Paris",
        }
    },
    "type": "customer_account_operations",
    "format": "csv",
    "options": {"useCashInventory": False, "groupPerMonth": False},
}

POLL_TIMEOUT_S = 90
POLL_DELAY_S = 3

PERIOD_HEADER = "From 18/06/2024 to 18/06/2024"
HEADER_ROW = 3  # lignes 1-2 = site + période fusionnées, ligne 3 = colonnes

SUMMARY_COLUMNS = ["Nom", "Prénom", "Société", "TTC", "HT", "TVA 10%", "Crédit", "Débit"]
DETAIL_COLUMNS = ["Nom", "Prénom", "Société", "Date", "Ticket #", "TTC", "HT", "TVA 10%"]
TICKET_DETAIL_COLUMNS = [
    "Nom", "Prénom", "Société", "Date", "Ticket #", "Contenu", "TTC", "HT", "TVA 10%",
]

CUSTOMER = "Cashpad NR auto test"

# (ticket, TTC, HT, TVA 10%) par opération liée à un ticket
DETAIL_ROWS = {
    570: (13.5, 12.27, 1.23),
    571: (-67.5, -61.38, -6.14),
    573: (-13.5, -12.27, -1.23),
}

# Lignes produit proratisées du 571 : 1 x British Salad (13.50) × ratio 2.5
PRORATED_LINE_571 = ("1 x British Salad", -33.75, -30.68, -3.07)

HT_ROUNDING_TOLERANCE = 0.021  # cf. en-tête : écart d'arrondi signalé, non figé


def get_env() -> dict:
    """Base staging + identifiants BO. Skip si absents : tout le monde n'a pas ce compte."""
    base = os.getenv("STAGING_BASE_URL")
    login = os.getenv("BOV2_STAGING_LOGIN")
    password = os.getenv("BOV2_STAGING_PASSWORD")

    missing = [
        name
        for name, value in (
            ("STAGING_BASE_URL", base),
            ("BOV2_STAGING_LOGIN", login),
            ("BOV2_STAGING_PASSWORD", password),
        )
        if not value
    ]
    if missing:
        pytest.skip(f"Variables manquantes dans .env : {missing}")

    host = urlparse(base).hostname
    assert host in HOST_ALLOWLIST, f"Host hors allowlist : {host!r} — refus d'exécuter"
    return {"base": base.rstrip("/"), "login": login, "password": password}


@pytest.fixture(scope="module")
def client() -> httpx.Client:
    """Client authentifié par JWT sso. Le token n'est jamais journalisé."""
    env = get_env()
    with httpx.Client(base_url=env["base"], timeout=60) as anon:
        response = anon.post(
            "/p/sso/public/1/sign-in",
            json={"username": env["login"], "password": env["password"]},
        )
        assert response.status_code in (200, 201), (
            f"sign-in a échoué : HTTP {response.status_code} — {response.text[:200]}"
        )
        token = response.json().get("token")
        assert token, f"sign-in n'a pas renvoyé de `token` : {sorted(response.json())}"

    with httpx.Client(
        base_url=env["base"],
        timeout=90,
        headers={"authorization": f"Bearer {token}", "app-timezone": "Europe/Paris"},
    ) as authed:
        yield authed


@pytest.fixture(scope="module")
def generated(client) -> dict:
    """Déclenche l'export, attend sa production, rend {exportId, entry, content}."""
    response = client.post(f"/p/exports/api/1/site/{SITE}/predefined-exports", json=EXPORT_BODY)
    assert response.status_code == 201, f"HTTP {response.status_code} : {response.text[:200]}"
    export_id = response.json()["exportId"]

    deadline = time.monotonic() + POLL_TIMEOUT_S
    entry = None
    while time.monotonic() < deadline:
        history = client.get(
            f"/p/exports/api/1/site/{SITE}/predefined-exports", params={"exportId": export_id}
        )
        assert history.status_code == 200, f"Historique : HTTP {history.status_code}"
        entry = next((e for e in history.json() if e.get("bucketKey")), None)
        if entry:
            break
        time.sleep(POLL_DELAY_S)
    assert entry, f"Export {export_id} toujours pas produit après {POLL_TIMEOUT_S}s"

    link = client.get(f"/p/exports/api/1/site/{SITE}/download-export/{entry['id']}")
    assert link.status_code == 200, f"Lien de téléchargement : HTTP {link.status_code}"
    url = link.json()
    storage_host = urlparse(url).hostname
    assert storage_host in STORAGE_HOST_ALLOWLIST, (
        f"Bucket hors allowlist : {storage_host!r} — refus de télécharger"
    )

    download = httpx.get(url, timeout=60)
    assert download.status_code == 200, f"Téléchargement : HTTP {download.status_code}"
    return {"exportId": export_id, "entry": entry, "content": download.content}


@pytest.fixture(scope="module")
def workbook(generated) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(generated["content"]))


def data_rows(ws) -> list[tuple]:
    """Lignes sous l'en-tête de colonnes, lignes vides exclues."""
    return [
        row
        for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True)
        if any(cell is not None for cell in row)
    ]


def columns_of(ws) -> list:
    return [c for c in next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))]


def test_01_export_is_an_xlsx_with_three_sheets(generated, workbook):
    """AC1 : l'export demandé pour un site et une période rend un xlsx à 3 onglets."""
    assert generated["entry"]["exportData"]["type"] == "customer_account_operations"
    assert generated["entry"]["bucketKey"].endswith(".xlsx"), generated["entry"]["bucketKey"]
    # signature ZIP : un xlsx, pas un CSV renommé (le front demande `csv`)
    assert generated["content"][:2] == b"PK", "Le fichier produit n'est pas un xlsx"
    assert len(workbook.worksheets) == 3, [ws.title for ws in workbook.worksheets]


def test_02_every_sheet_has_the_site_and_period_header(workbook):
    """AC6 (en-têtes) : site + période fusionnés en A1 / A2 sur chaque onglet.

    C'était le ❌ du 2026-09-11 (en-tête absent des 3 onglets) : ce step garde le
    correctif.
    """
    for ws in workbook.worksheets:
        merged = {str(r) for r in ws.merged_cells.ranges}
        last = ws.cell(row=HEADER_ROW, column=ws.max_column).column_letter
        assert f"A1:{last}1" in merged and f"A2:{last}2" in merged, (
            f"{ws.title!r} : en-tête non fusionné ({sorted(merged)})"
        )
        assert ws["A1"].value, f"{ws.title!r} : nom du site absent en A1"
        assert ws["A2"].value == PERIOD_HEADER, f"{ws.title!r} : A2 = {ws['A2'].value!r}"


def test_03_columns_match_the_bov1_structure(workbook):
    """AC6 (colonnes) : Résumé / Détail / Détail tickets, une colonne par taux de TVA."""
    summary, detail, tickets = workbook.worksheets
    assert columns_of(summary) == SUMMARY_COLUMNS
    assert columns_of(detail) == DETAIL_COLUMNS
    assert columns_of(tickets) == TICKET_DETAIL_COLUMNS


def test_04_summary_totals_are_the_sum_of_the_operations(workbook):
    """AC2 : les totaux du Résumé égalent la somme des opérations du Détail."""
    summary, detail, _ = workbook.worksheets
    rows = data_rows(summary)
    assert [r[0] for r in rows] == [CUSTOMER], f"Clients du Résumé : {[r[0] for r in rows]}"
    _, _, _, ttc, ht, vat, credit, debit = rows[0]

    ops = data_rows(detail)
    assert ttc == pytest.approx(sum(r[5] for r in ops), abs=0.005)
    assert ht == pytest.approx(sum(r[6] for r in ops), abs=0.005)
    assert vat == pytest.approx(sum(r[7] for r in ops), abs=0.005)

    assert (ttc, vat, credit, debit) == pytest.approx((-67.5, -6.14, 200, 0), abs=0.005)
    assert ht == pytest.approx(-61.38, abs=HT_ROUNDING_TOLERANCE)


def test_05_partial_receipt_payment_is_prorated(workbook):
    """AC3 : TTC / HT / TVA d'une opération liée à un ticket sont proratisés.

    Ticket 571 : opération 67.50 sur un ticket de 27.00 → ratio 2.5, appliqué
    à l'opération ET à chaque ligne produit du Détail tickets.
    """
    _, detail, tickets = workbook.worksheets

    by_ticket = {r[4]: r for r in data_rows(detail)}
    assert set(by_ticket) == set(DETAIL_ROWS), f"Tickets du Détail : {sorted(by_ticket)}"
    for ticket, (ttc, ht, vat) in DETAIL_ROWS.items():
        row = by_ticket[ticket]
        assert row[0] == CUSTOMER
        assert (row[5], row[7]) == pytest.approx((ttc, vat), abs=0.005), f"ticket {ticket}"
        assert row[6] == pytest.approx(ht, abs=HT_ROUNDING_TOLERANCE), f"ticket {ticket}"

    rows = data_rows(tickets)
    start = next(i for i, r in enumerate(rows) if r[4] == 571)
    lines = []
    for r in rows[start + 1:]:
        if r[4] is not None:  # ligne d'opération suivante
            break
        lines.append(r[5:])
    assert len(lines) == 2, f"Lignes produit du 571 : {lines}"
    for content, ttc, ht, vat in lines:
        assert content == PRORATED_LINE_571[0]
        assert (ttc, ht, vat) == pytest.approx(PRORATED_LINE_571[1:], abs=0.005)

    # la somme des lignes produit redonne le TTC et la TVA de l'opération
    assert sum(l[1] for l in lines) == pytest.approx(by_ticket[571][5], abs=0.005)
    assert sum(l[3] for l in lines) == pytest.approx(by_ticket[571][7], abs=0.005)


def test_06_receipt_less_operation_is_only_in_credit_debit(workbook):
    """AC4 : l'opération sans ticket (+200) ne vit que dans Crédit / Débit.

    AC5 (montant nul) : aucune ligne à 0 — non discriminant sur cette donnée,
    qui ne contient aucune opération nulle (cf. en-tête).
    """
    summary, detail, tickets = workbook.worksheets
    assert data_rows(summary)[0][6] == pytest.approx(200)

    for ws in (detail, tickets):
        ops = [r for r in data_rows(ws) if r[0] is not None]
        assert all(r[4] is not None for r in ops), f"{ws.title!r} : opération sans ticket"
        assert all(r[-3] != 0 for r in ops), f"{ws.title!r} : opération à montant nul"
    assert 200 not in [r[5] for r in data_rows(detail)]
